"""M7: pengawasan (akses, monitor, WebSocket, tindakan), pengumuman, hasil & ekspor."""

import uuid
import zipfile
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from starlette.websockets import WebSocketDisconnect

from app import proctor
from app.sweeper import sweep_once
from tests.conftest import ORIGIN, login, make_user
from tests.test_attempts import PUBLISH, _sql, exam, exam_body, idem, start  # noqa: F401
from tests.test_exams import _now, setup  # noqa: F401
from tests.test_questions import PG

FORBIDDEN = ("answer_key", "is_correct", "correct_option_ids", "score", '"answer"')


def ws_headers(client: TestClient, origin: str = ORIGIN) -> dict[str, str]:
    token = next(c.value for c in client.cookies.jar if c.name == "localcbt_session")
    return {"origin": origin, "cookie": f"localcbt_session={token}"}


@pytest.fixture
def pengawas(exam: dict[str, Any]) -> TestClient:
    make_user("pengawas", "PROCTOR", name="Pak Pengawas")
    return login("pengawas")


def _assign(exam: dict[str, Any], username: str) -> str:
    uid = str(_sql("select id from users where username = :u", u=username).scalar())
    r = exam["client"].post(f"/api/v1/exams/{exam['exam']}/proctors", json={"user_id": uid})
    assert r.status_code == 201, r.text
    return uid


# ---- Akses (T-021) -------------------------------------------------------------


def test_t021_unassigned_proctor_denied_and_audited(
    exam: dict[str, Any], pengawas: TestClient
) -> None:
    attempt_id = start(exam["student"], exam["exam"]).json()["id"]
    deadline = _sql("select effective_deadline_at from exam_attempts").scalar()
    base = f"/api/v1/proctor/attempts/{attempt_id}"
    reason = {"reason": "Alasan cukup panjang"}
    assert pengawas.get("/api/v1/proctor/exams").json() == []
    assert pengawas.get(f"/api/v1/proctor/exams/{exam['exam']}/monitor").status_code == 404
    assert (
        pengawas.post(f"{base}/time-extensions", json={**reason, "seconds": 600}).status_code == 404
    )
    assert pengawas.post(f"{base}/approve-device-switch", json=reason).status_code == 404
    assert pengawas.get(f"{base}/events").status_code == 404
    r = pengawas.post(f"/api/v1/proctor/exams/{exam['exam']}/announcements", json={"message": "x"})
    assert r.status_code == 404
    assert _sql("select effective_deadline_at from exam_attempts").scalar() == deadline
    assert (
        _sql("select count(*) from audit_events where action = 'PROCTOR_ACTION_DENIED'").scalar()
        == 3
    )

    _assign(exam, "pengawas")
    assert [e["title"] for e in pengawas.get("/api/v1/proctor/exams").json()] == ["PTS Matematika"]
    assert (
        pengawas.post(f"{base}/time-extensions", json={**reason, "seconds": 600}).status_code == 201
    )

    make_user("guru.b", "TEACHER")
    other = login("guru.b")
    uid = _sql("select id from users where username = 'guru.b'").scalar()
    r = other.post(f"/api/v1/exams/{exam['exam']}/proctors", json={"user_id": str(uid)})
    assert r.status_code == 404  # hanya pemilik ujian yang menugaskan pengawas


def test_websocket_streams_snapshots_and_closes_when_revoked(
    exam: dict[str, Any], pengawas: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(proctor, "MONITOR_PUSH_SECONDS", 0.05)
    uid = _assign(exam, "pengawas")
    start(exam["student"], exam["exam"])
    url = f"/api/v1/ws/exams/{exam['exam']}"
    with pengawas.websocket_connect(url, headers=ws_headers(pengawas)) as ws:
        first = ws.receive_json()
        assert first["type"] == "snapshot" and first["counts"]["active"] == 1
        assert not any(k in str(first) for k in FORBIDDEN)
        exam["client"].delete(f"/api/v1/exams/{exam['exam']}/proctors/{uid}")
        with pytest.raises(WebSocketDisconnect) as closed:
            for _ in range(50):
                ws.receive_json()
        assert closed.value.code == 1008

    with (
        pytest.raises(WebSocketDisconnect),
        pengawas.websocket_connect(url, headers=ws_headers(exam["client"], "https://evil.example")),
    ):
        pass
    with (
        pytest.raises(WebSocketDisconnect),
        pengawas.websocket_connect(url, headers={"origin": ORIGIN}),
    ):
        pass


def test_t024_two_teachers_two_live_exams_no_leakage(exam: dict[str, Any]) -> None:
    s = exam
    start(s["student"], s["exam"])
    teacher_b = make_user("guru.b", "TEACHER")
    _sql(
        "insert into teaching_assignments (teacher_id, class_id, subject_id) values (:t, :c, :s)",
        t=teacher_b,
        c=s["class_b"],
        s=s["subject"],
    )
    b = login("guru.b")
    bank = b.post("/api/v1/question-banks", json={"title": "B", "subject_id": s["subject"]}).json()[
        "id"
    ]
    qid = b.post(f"/api/v1/question-banks/{bank}/questions", json=PG).json()["id"]
    body = exam_body(
        s,
        title="Ujian B",
        class_ids=[s["class_b"]],
        question_ids=[qid],
        start_window_from=_now().isoformat(),
    )
    exam_b = b.post("/api/v1/exams", json=body).json()["id"]
    assert b.post(f"/api/v1/exams/{exam_b}/publish", json=PUBLISH).status_code == 200
    student_b = make_user("0055555", "STUDENT", name="Siswa B")
    _sql(
        "insert into class_memberships (class_id, student_id) values (:c, :s)",
        c=s["class_b"],
        s=student_b,
    )

    mon_a = s["client"].get(f"/api/v1/proctor/exams/{s['exam']}/monitor").json()
    mon_b = b.get(f"/api/v1/proctor/exams/{exam_b}/monitor").json()
    assert [p["nis"] for p in mon_a["participants"]] == ["0012345"]
    assert [p["nis"] for p in mon_b["participants"]] == ["0055555"]
    assert s["client"].get(f"/api/v1/proctor/exams/{exam_b}/monitor").status_code == 404
    assert b.get(f"/api/v1/teacher/exams/{s['exam']}/results").status_code == 404
    assert b.get(f"/api/v1/teacher/exams/{s['exam']}/results/export").status_code == 404


# ---- Monitor -------------------------------------------------------------------


def test_roster_statuses_and_no_answers_or_scores(exam: dict[str, Any]) -> None:
    s = exam
    # Siswa di dua kelas yang sama-sama ditugaskan muncul sekali.
    _sql(
        "insert into exam_assignments (exam_id, class_id) values (:e, :c)",
        e=s["exam"],
        c=s["class_b"],
    )
    _sql(
        "insert into class_memberships (class_id, student_id) values (:c, :s)",
        c=s["class_b"],
        s=s["student_id"],
    )
    names = {
        "0010001": "NOT_STARTED",
        "0010002": "DISCONNECTED",
        "0010003": "PAUSED",
        "0010004": "SUBMITTED",
    }
    for nis in names:
        sid = make_user(nis, "STUDENT")
        _sql(
            "insert into class_memberships (class_id, student_id) values (:c, :s)",
            c=s["class"],
            s=sid,
        )
    att = start(s["student"], s["exam"]).json()
    q = att["questions"][0]
    s["student"].put(
        f"/api/v1/student/attempts/{att['id']}/answers/{q['id']}",
        json={"answer": None, "client_revision": 1, "marked_for_review": True},
        headers=idem(),
    )
    for nis in ("0010002", "0010003", "0010004"):
        c = login(nis)
        a = start(c, s["exam"]).json()
        if nis == "0010004":
            c.post(
                f"/api/v1/student/attempts/{a['id']}/submit", json={"confirm": True}, headers=idem()
            )
    _sql(
        "update exam_attempts a set last_heartbeat_at = now() - interval '20 seconds' "
        "from users u where u.id = a.student_id and u.username in ('0010002', '0010003')"
    )
    _sql(
        "update exam_attempts a set state = 'PAUSED', paused_at = now() "
        "from users u where u.id = a.student_id and u.username = '0010003'"
    )
    # Siswa yang sudah mulai lalu dipindah kelas tetap terpantau.
    _sql(
        "delete from class_memberships m using users u "
        "where u.id = m.student_id and u.username = '0010002'"
    )

    r = s["client"].get(f"/api/v1/proctor/exams/{s['exam']}/monitor")
    rows = {p["nis"]: p for p in r.json()["participants"]}
    assert len(r.json()["participants"]) == 5
    assert {nis: rows[nis]["status"] for nis in names} == names
    assert rows["0012345"]["status"] == "ACTIVE" and rows["0012345"]["answered_count"] == 0
    assert r.json()["counts"] == {
        "not_started": 1,
        "active": 1,
        "disconnected": 1,
        "paused": 1,
        "submitted": 1,
    }
    assert not any(k in r.text for k in FORBIDDEN)


def test_t022_tab_hidden_is_counted_per_transition_only(exam: dict[str, Any]) -> None:
    att = start(exam["student"], exam["exam"]).json()
    url = f"/api/v1/student/attempts/{att['id']}/heartbeat"
    for visible in (True, False, False, True, False):
        exam["student"].post(url, json={"tab_visible": visible})
    assert _sql("select count(*) from audit_events where action = 'TAB_HIDDEN'").scalar() == 2
    mon = exam["client"].get(f"/api/v1/proctor/exams/{exam['exam']}/monitor").json()
    assert mon["participants"][0]["tab_hidden"] == 2
    assert _sql("select state from exam_attempts").scalar() == "IN_PROGRESS"
    events = exam["client"].get(f"/api/v1/proctor/attempts/{att['id']}/events").json()
    assert [e["kind"] for e in events].count("TAB_HIDDEN") == 2


# ---- Tindakan ------------------------------------------------------------------


def test_extend_time_moves_deadline_and_is_audited(exam: dict[str, Any]) -> None:
    att = start(exam["student"], exam["exam"]).json()
    url = f"/api/v1/proctor/attempts/{att['id']}/time-extensions"
    before = _sql("select effective_deadline_at from exam_attempts").scalar()
    for bad in (59, 7201):
        assert (
            exam["client"]
            .post(url, json={"seconds": bad, "reason": "Alasan cukup panjang"})
            .status_code
            == 422
        )
    assert exam["client"].post(url, json={"seconds": 600}).status_code == 422  # alasan wajib
    r = exam["client"].post(url, json={"seconds": 600, "reason": "Listrik padam di ruang 3"})
    assert r.status_code == 201
    after = _sql("select effective_deadline_at from exam_attempts").scalar()
    assert (after - before).total_seconds() == 600
    row = _sql("select reason, metadata from audit_events where action = 'TIME_EXTENDED'").one()
    assert row.reason == "Listrik padam di ruang 3" and row.metadata == {"seconds": 600}

    # Batas lewat tapi sweeper belum jalan: tidak boleh menghidupkan attempt.
    _sql("update exam_attempts set effective_deadline_at = now() - interval '1 second'")
    r = exam["client"].post(url, json={"seconds": 600, "reason": "Terlambat diberi waktu"})
    assert r.status_code == 409
    assert _sql("select state from exam_attempts").scalar() == "GRADED"


def test_review_stays_closed_while_any_student_still_works(exam: dict[str, Any]) -> None:
    s = exam
    a = start(s["student"], s["exam"]).json()
    s["student"].post(
        f"/api/v1/student/attempts/{a['id']}/submit", json={"confirm": True}, headers=idem()
    )
    other = make_user("0012346", "STUDENT")
    _sql(
        "insert into class_memberships (class_id, student_id) values (:c, :s)",
        c=s["class"],
        s=other,
    )
    b = start(login("0012346"), s["exam"]).json()
    s["client"].post(
        f"/api/v1/proctor/attempts/{b['id']}/time-extensions",
        json={"seconds": 7200, "reason": "Kebutuhan khusus siswa"},
    )
    _sql(
        'update exams set review_policy = \'{"enabled": true, "show_correct_answers": true}\', '
        "start_window_from = now() - interval '3 hours', "
        "start_window_until = now() - interval '2 hours', duration_sec = 60"
    )
    review = f"/api/v1/student/attempts/{a['id']}/review"
    assert s["student"].get(review).status_code == 403  # siswa lain masih mengerjakan
    _sql(
        "update exam_attempts set effective_deadline_at = now() - interval '1 second' "
        "where id = :i",
        i=b["id"],
    )
    sweep_once()
    assert s["student"].get(review).status_code == 200


def test_announcements_persisted_and_scoped(exam: dict[str, Any]) -> None:
    att = start(exam["student"], exam["exam"]).json()
    r = exam["client"].post(
        f"/api/v1/proctor/exams/{exam['exam']}/announcements",
        json={"message": "Sisa waktu 30 menit, periksa kembali jawaban."},
    )
    assert r.status_code == 201
    beat = exam["student"].post(f"/api/v1/student/attempts/{att['id']}/heartbeat", json={}).json()
    assert beat["announcements_count"] == 1
    listed = exam["student"].get(f"/api/v1/student/attempts/{att['id']}/announcements").json()
    assert [a["message"] for a in listed] == ["Sisa waktu 30 menit, periksa kembali jawaban."]
    make_user("0099999", "STUDENT")
    assert (
        login("0099999").get(f"/api/v1/student/attempts/{att['id']}/announcements").status_code
        == 404
    )


# ---- Hasil ---------------------------------------------------------------------


def test_publish_unpublish_controls_student_result_not_review(exam: dict[str, Any]) -> None:
    s = exam
    att = start(s["student"], s["exam"]).json()
    s["student"].post(
        f"/api/v1/student/attempts/{att['id']}/submit", json={"confirm": True}, headers=idem()
    )
    s["client"].put(
        f"/api/v1/teacher/exams/{s['exam']}/result-policy", json={"visibility": "AFTER_PUBLICATION"}
    )
    result = f"/api/v1/student/attempts/{att['id']}/result"
    assert s["student"].get(result).status_code == 403

    rows = s["client"].get(f"/api/v1/teacher/exams/{s['exam']}/results").json()["rows"]
    assert rows[0]["status"] == "GRADED" and rows[0]["max_score"] == 7.5
    assert s["client"].post(f"/api/v1/teacher/exams/{s['exam']}/results/publish").status_code == 200
    assert s["student"].get(result).status_code == 200
    assert s["student"].get(f"/api/v1/student/attempts/{att['id']}/review").status_code == 403
    s["client"].post(f"/api/v1/teacher/exams/{s['exam']}/results/unpublish")
    assert s["student"].get(result).status_code == 403
    actions = [r[0] for r in _sql("select action from audit_events order by created_at")]
    assert {"RESULT_POLICY_CHANGED", "RESULTS_PUBLISHED", "RESULTS_UNPUBLISHED"} <= set(actions)


def test_regrade_only_touches_ungraded(exam: dict[str, Any]) -> None:
    s = exam
    a = start(s["student"], s["exam"]).json()
    other = make_user("0012346", "STUDENT")
    _sql(
        "insert into class_memberships (class_id, student_id) values (:c, :s)",
        c=s["class"],
        s=other,
    )
    start(login("0012346"), s["exam"])
    s["student"].post(
        f"/api/v1/student/attempts/{a['id']}/submit", json={"confirm": True}, headers=idem()
    )
    graded_at = _sql("select graded_at from exam_attempts where id = :i", i=a["id"]).scalar()
    # Simulasikan penilaian gagal: SUBMITTED tanpa skor.
    _sql(
        "update exam_attempts set state = 'SUBMITTED', submitted_at = now(), "
        "submit_reason = 'STUDENT' "
        "where student_id = :s",
        s=other,
    )
    r = s["client"].post(f"/api/v1/teacher/exams/{s['exam']}/regrade")
    assert r.json() == {"pending": 1, "graded": 1}
    assert _sql("select count(*) from exam_attempts where state = 'GRADED'").scalar() == 2
    assert (
        _sql("select graded_at from exam_attempts where id = :i", i=a["id"]).scalar() == graded_at
    )


def test_export_matches_db_and_neutralizes_formulas(exam: dict[str, Any]) -> None:
    s = exam
    _sql(
        'update users set display_name = \'=HYPERLINK("http://x","klik")\' '
        "where username = '0012345'"
    )
    att = start(s["student"], s["exam"]).json()
    q = att["questions"][0]
    opt = q["options"][0]["id"] if q["options"] else None
    if opt:
        s["student"].put(
            f"/api/v1/student/attempts/{att['id']}/answers/{q['id']}",
            json={
                "answer": {"type": q["type"], "selected_option_id": opt}
                if q["type"] == "SINGLE_CHOICE"
                else {"type": q["type"], "selected_option_ids": [opt]},
                "client_revision": 1,
                "marked_for_review": False,
            },
            headers=idem(),
        )
    s["student"].post(
        f"/api/v1/student/attempts/{att['id']}/submit", json={"confirm": True}, headers=idem()
    )

    r = s["client"].get(f"/api/v1/teacher/exams/{s['exam']}/results/export")
    assert r.status_code == 200
    assert b"<f>" not in zipfile.ZipFile(BytesIO(r.content)).read("xl/worksheets/sheet1.xml")
    ws = load_workbook(BytesIO(r.content)).active
    assert ws is not None
    header, row = [c.value for c in ws[1]], [c.value for c in ws[2]]
    assert header[:6] == ["NIS", "Nama", "Kelas", "Status", "Nilai", "Maks"]
    assert row[1] == '=HYPERLINK("http://x","klik")' and row[3] == "Dinilai"
    total, maximum = _sql("select score_total, score_max from exam_attempts").one()
    assert Decimal(str(row[4])) == total and Decimal(str(row[5])) == maximum
    per_q = {
        f"S{o}": e
        for o, e in _sql(
            "select x.ordinal, q.earned from question_scores q "
            "join exam_question_snapshots x on x.id = q.snapshot_question_id"
        )
    }
    for col, value in zip(header[6:], row[6:], strict=True):
        assert Decimal(str(value)) == per_q[col]
    assert uuid.UUID(str(s["exam"]))
