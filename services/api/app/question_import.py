"""Import soal dari XLSX (satu baris = satu soal) atau DOCX (satu tabel label|isi = satu soal).

Keduanya menghasilkan dict {label HEADER: teks} lalu divalidasi QuestionIn yang sama dengan editor.
Gambar tidak diimport (ditambahkan lewat editor) agar tidak ada konten yang hilang diam-diam.
"""

import hashlib
import json
from io import BytesIO
from typing import Annotated, Any, Literal
from uuid import UUID
from xml.etree.ElementTree import Element, ParseError
from xml.sax.saxutils import escape

from defusedxml import DefusedXmlException
from defusedxml.ElementTree import fromstring
from fastapi import APIRouter, Depends, Response, UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy import text

from app.admin import MAX_UPLOAD, cell_text, check_zip
from app.auth import CurrentUser, require_roles
from app.db import DB
from app.errors import ApiError, not_found
from app.questions import QuestionIn, get_bank, save_question

router = APIRouter(prefix="/api/v1", tags=["Question Bank"])
Teacher = Annotated[CurrentUser, Depends(require_roles("TEACHER"))]

LETTERS = "ABCDE"
HEADER = ["Jenis", "Stimulus", "Pertanyaan", "Bobot", "Kunci"] + [f"Opsi {c}" for c in LETTERS]
TYPES = {"PG": "SINGLE_CHOICE", "MCMA": "MULTIPLE_CHOICE", "KATEGORI": "CATEGORY_CHOICE"}
MAX_QUESTIONS = 500
EXAMPLES = [
    {
        "Jenis": "PG",
        "Pertanyaan": "Hasil dari 12 × 3 adalah …",
        "Bobot": "1",
        "Kunci": "B",
        "Opsi A": "32",
        "Opsi B": "36",
        "Opsi C": "38",
        "Opsi D": "42",
    },
    {
        "Jenis": "MCMA",
        "Stimulus": "Perhatikan bilangan 2, 3, 4, 9.",
        "Pertanyaan": "Manakah yang merupakan bilangan prima?",
        "Bobot": "2",
        "Kunci": "A,B",
        "Opsi A": "2",
        "Opsi B": "3",
        "Opsi C": "4",
        "Opsi D": "9",
    },
    {
        "Jenis": "KATEGORI",
        "Stimulus": "Perhatikan bilangan 2, 3, 4, 9.",
        "Pertanyaan": "Tentukan Benar (B) atau Salah (S) untuk setiap pernyataan.",
        "Bobot": "3",
        "Kunci": "B,S,B",
        "Opsi A": "4 adalah bilangan genap",
        "Opsi B": "9 adalah bilangan prima",
        "Opsi C": "2 + 3 = 5",
    },
]

Row = tuple[int, dict[str, str], list[tuple[str, str]]]  # (nomor, sel, error parse)
FIELD = {"stem": "Pertanyaan", "weight": "Bobot", "options": "Opsi", "statements": "Opsi"}


def to_question(cells: dict[str, str]) -> tuple[dict[str, Any] | None, list[tuple[str, str]]]:
    """Sel template → payload QuestionIn, atau daftar (kolom, pesan)."""
    qtype = TYPES.get(cells.get("Jenis", "").strip().upper())
    if qtype is None:
        return None, [("Jenis", "Isi dengan PG, MCMA, atau KATEGORI.")]
    filled = [(c, cells[f"Opsi {c}"]) for c in LETTERS if cells.get(f"Opsi {c}", "").strip()]
    keys = [k.strip().upper() for k in cells.get("Kunci", "").split(",") if k.strip()]
    payload: dict[str, Any] = {
        "type": qtype,
        "stem": cells.get("Pertanyaan", ""),
        "weight": cells.get("Bobot", "").replace(",", ".") or "1",
    }
    if qtype == "CATEGORY_CHOICE":
        if len(keys) != len(filled) or not set(keys) <= {"B", "S"}:
            return None, [("Kunci", "Isi B/S untuk setiap pernyataan, mis. B,S,B.")]
        payload["statements"] = [
            {"text": t, "correct_value": k == "B"} for (_, t), k in zip(filled, keys, strict=True)
        ]
    else:
        if not keys or not set(keys) <= {c for c, _ in filled}:
            return None, [("Kunci", "Kunci harus huruf opsi yang terisi, mis. B atau A,C.")]
        payload["options"] = [{"text": t, "is_correct": c in keys} for c, t in filled]
    try:
        QuestionIn.model_validate(payload)
    except ValidationError as e:
        return None, [
            (
                FIELD.get(str(err["loc"][0]), "Soal") if err["loc"] else "Soal",
                err["msg"].removeprefix("Value error, "),
            )
            for err in e.errors()
        ]
    return payload, []


# ---- XLSX ------------------------------------------------------------------------


def parse_xlsx(data: bytes) -> list[Row]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [cell_text(c).lower() for c in next(rows, ())[: len(HEADER)]]
        if header != [h.lower() for h in HEADER]:
            raise ApiError(
                422, "INVALID_TEMPLATE", "Kolom tidak sesuai template. Unduh template terbaru."
            )
        out: list[Row] = []
        for number, row in enumerate(rows, start=2):
            values = [cell_text(c) for c in row[: len(HEADER)]]
            if not any(values):
                continue
            out.append((number, dict(zip(HEADER, values, strict=False)), []))
        wb.close()
        return out
    except ApiError:
        raise
    except Exception as e:  # berkas tak tepercaya: semua kegagalan parse = berkas tidak valid
        raise ApiError(422, "INVALID_FILE", "Berkas XLSX tidak dapat dibaca.") from e


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Soal"
    ws.append(HEADER)
    for ex in EXAMPLES:
        ws.append([ex.get(h, "") for h in HEADER])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- DOCX ------------------------------------------------------------------------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_LABELS = {h.lower(): h for h in HEADER}


def _cell_text(tc: Element) -> str:
    # Word memecah teks ke banyak run; gabungkan semua w:t per paragraf.
    paragraphs = []
    for p in tc.iter(f"{W}p"):
        parts = []
        for el in p.iter():
            if el.tag == f"{W}t":
                parts.append(el.text or "")
            elif el.tag == f"{W}tab":
                parts.append("\t")
            elif el.tag in (f"{W}br", f"{W}cr"):
                parts.append("\n")
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs).strip()


def parse_docx(data: bytes) -> list[Row]:
    zf = check_zip(data)
    try:
        root = fromstring(zf.read("word/document.xml"))
    except KeyError as e:
        raise ApiError(422, "INVALID_FILE", "Berkas bukan dokumen Word (.docx).") from e
    except (ParseError, DefusedXmlException) as e:  # termasuk DTD/entity bomb
        raise ApiError(422, "INVALID_FILE", "Isi dokumen Word tidak dapat dibaca.") from e
    body = root.find(f"{W}body")
    tables = body.findall(f"{W}tbl") if body is not None else []
    if not tables:
        raise ApiError(
            422,
            "INVALID_TEMPLATE",
            "Tidak ada tabel soal. Gunakan template DOCX (1 tabel = 1 soal).",
        )
    out: list[Row] = []
    for number, tbl in enumerate(tables, start=1):
        cells: dict[str, str] = {}
        errors: list[tuple[str, str]] = []
        for tr in tbl.findall(f"{W}tr"):
            tcs = tr.findall(f"{W}tc")
            if len(tcs) != 2:
                errors.append(("Tabel", "Setiap baris tabel harus 2 kolom: label | isi."))
                continue
            label = _cell_text(tcs[0]).rstrip(":").strip()
            if label.lower() not in _LABELS:
                errors.append(("Tabel", f"Label tidak dikenal: {label[:40]}"))
                continue
            value = tcs[1]
            if any(
                next(value.iter(f"{W}{tag}"), None) is not None
                for tag in ("tbl", "drawing", "pict", "object")
            ):
                errors.append(
                    (
                        _LABELS[label.lower()],
                        "Gambar/tabel di dalam sel belum didukung lewat import.",
                    )
                )
                continue
            cells[_LABELS[label.lower()]] = _cell_text(value)
        out.append((number, cells, errors))
    return out


def build_docx(tables: list[list[tuple[str, str]]]) -> bytes:
    """DOCX minimal (tanpa dependensi): satu tabel 2 kolom per soal."""
    import zipfile

    def cell(t: str) -> str:
        return f'<w:tc><w:p><w:r><w:t xml:space="preserve">{escape(t)}</w:t></w:r></w:p></w:tc>'

    border = "".join(
        f'<w:{side} w:val="single" w:sz="4" w:color="999999"/>'
        for side in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    body = "".join(
        f"<w:tbl><w:tblPr><w:tblBorders>{border}</w:tblBorders></w:tblPr>"
        + "".join(f"<w:tr>{cell(k)}{cell(v)}</w:tr>" for k, v in rows)
        + "</w:tbl><w:p/>"
        for rows in tables
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        z.writestr("word/document.xml", document)
    return buf.getvalue()


def template_docx() -> bytes:
    return build_docx([[(h, ex[h]) for h in HEADER if ex.get(h)] for ex in EXAMPLES])


# ---- Endpoint --------------------------------------------------------------------

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@router.get("/question-import-templates/{fmt}")
def download_template(fmt: Literal["xlsx", "docx"], _: Teacher) -> Response:
    data, mime = (template_xlsx(), XLSX_MIME) if fmt == "xlsx" else (template_docx(), DOCX_MIME)
    return Response(
        data,
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="template-soal.{fmt}"'},
    )


@router.post("/question-banks/{bank_id}/imports/preview")
def preview_question_import(
    bank_id: UUID, file: UploadFile, user: Teacher, db: DB
) -> dict[str, Any]:
    get_bank(db, user, bank_id, "EDIT")
    data = file.file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise ApiError(413, "FILE_TOO_LARGE", "Ukuran berkas maksimal 2 MB.")
    names = set(check_zip(data).namelist())
    if "word/document.xml" in names:
        kind, rows = "QUESTION_DOCX", parse_docx(data)
    elif "xl/workbook.xml" in names:
        kind, rows = "QUESTION_XLSX", parse_xlsx(data)
    else:
        raise ApiError(422, "INVALID_FILE", "Gunakan berkas .xlsx atau .docx dari template.")
    if not rows:
        raise ApiError(422, "EMPTY_FILE", "Berkas tidak berisi soal.")
    if len(rows) > MAX_QUESTIONS:
        raise ApiError(413, "TOO_MANY_ROWS", f"Maksimal {MAX_QUESTIONS} soal per berkas.")

    errors: list[dict[str, Any]] = []
    valid: list[dict[str, Any]] = []
    for number, cells, parse_errors in rows:
        payload, row_errors = (None, parse_errors) if parse_errors else to_question(cells)
        errors += [{"row_number": number, "field": f, "message": m} for f, m in row_errors]
        if payload is not None:
            valid.append({"question": payload, "stimulus": cells.get("Stimulus", "").strip()})

    db.execute(
        text(
            "delete from import_jobs where status = 'PREVIEW' "
            "and created_at < now() - interval '1 hour'"
        )
    )
    job_id = db.scalar(
        text("""
        insert into import_jobs (type, created_by, bank_id, status, source_sha256, rows, errors)
        values (:type, :by, :b, 'PREVIEW', :sha, cast(:rows as jsonb), cast(:errors as jsonb))
        returning id
        """),
        {
            "type": kind,
            "by": user.id,
            "b": bank_id,
            "sha": hashlib.sha256(data).hexdigest(),
            "rows": json.dumps(valid),
            "errors": json.dumps(errors),
        },
    )
    return {
        "preview_id": str(job_id),
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "errors": errors,
    }


class ImportCommitIn(BaseModel):
    preview_id: UUID
    confirm: Literal[True]


@router.post("/question-banks/{bank_id}/imports/commit")
def commit_question_import(
    bank_id: UUID, body: ImportCommitIn, user: Teacher, db: DB
) -> dict[str, Any]:
    get_bank(db, user, bank_id, "EDIT")
    job = (
        db.execute(
            text("""
            select status, rows, errors from import_jobs
            where id = :id and bank_id = :b and created_by = :by
              and type in ('QUESTION_XLSX', 'QUESTION_DOCX')
            for update
            """),
            {"id": body.preview_id, "b": bank_id, "by": user.id},
        )
        .mappings()
        .first()
    )
    if job is None:
        raise not_found()
    if job["status"] != "PREVIEW":
        raise ApiError(409, "IMPORT_ALREADY_COMMITTED", "Import ini sudah disimpan.")
    if job["errors"]:
        raise ApiError(
            409, "IMPORT_HAS_ERRORS", "Perbaiki semua soal bermasalah lalu unggah ulang berkas."
        )
    # Stimulus dengan teks sama dalam satu berkas = satu stimulus bersama (pengelompokan soal).
    stimuli: dict[str, UUID] = {}
    for row in job["rows"]:
        if row["stimulus"] and row["stimulus"] not in stimuli:
            sid = db.scalar(
                text("insert into stimuli (bank_id, body_text) values (:b, :t) returning id"),
                {"b": bank_id, "t": row["stimulus"]},
            )
            assert sid is not None
            stimuli[row["stimulus"]] = sid
        question = QuestionIn.model_validate(
            {**row["question"], "stimulus_id": stimuli.get(row["stimulus"])}
        )
        save_question(db, bank_id, question, user.id)
    db.execute(
        text("update import_jobs set status = 'COMPLETED', rows = null where id = :id"),
        {"id": body.preview_id},
    )
    return {"job_id": str(body.preview_id), "status": "COMPLETED", "imported": len(job["rows"])}
