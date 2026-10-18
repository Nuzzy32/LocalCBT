"""Administrasi akademik: kelas, mapel, staf, penugasan mengajar, akun & import siswa."""

import hashlib
import json
import re
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from typing import Annotated, Any, Literal
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app.auth import CurrentUser, generate_password, hash_password, require_roles
from app.db import DB
from app.errors import ApiError, not_found

router = APIRouter(prefix="/api/v1/admin", tags=["Academic Admin"])
Admin = Annotated[CurrentUser, Depends(require_roles("ADMIN"))]
AdminOrTeacher = Annotated[CurrentUser, Depends(require_roles("ADMIN", "TEACHER"))]

USERNAME = r"^[0-9A-Za-z._-]{1,64}$"
NIS_RE = re.compile(r"^[0-9A-Za-z.-]{1,32}$")
MAX_UPLOAD = 2 * 1024 * 1024
MAX_UNZIPPED = 20 * 1024 * 1024
MAX_ROWS = 2000


def _hash_many(passwords: list[str]) -> list[str]:
    # Argon2 melepas GIL; 4 thread × 64 MiB memory_cost = ~256 MiB puncak.
    with ThreadPoolExecutor(max_workers=4) as pool:
        return list(pool.map(hash_password, passwords))


# ---- Kelas & mapel ---------------------------------------------------------------


class ClassIn(BaseModel):
    code: str = Field(min_length=1, max_length=32, pattern=r"^[0-9A-Za-z ._-]+$")
    label: str = Field(min_length=1, max_length=160)
    academic_year: str = Field(pattern=r"^\d{4}/\d{4}$")


@router.get("/classes")
def list_classes(user: AdminOrTeacher, db: DB) -> list[dict[str, Any]]:
    # Guru hanya melihat kelas yang ia ajar (object-level, bukan sekadar role).
    rows = db.execute(
        text("""
        select c.id, c.code, c.label, c.academic_year, c.active
        from classes c
        where :admin or exists (
          select 1 from teaching_assignments t where t.class_id = c.id and t.teacher_id = :uid)
        order by c.academic_year desc, c.code
        """),
        {"admin": user.has("ADMIN"), "uid": user.id},
    ).mappings()
    return [{**r, "id": str(r["id"]), "name": r["label"]} for r in rows]


@router.post("/classes", status_code=201)
def create_class(body: ClassIn, _: Admin, db: DB) -> dict[str, Any]:
    class_id = db.scalar(
        text("""
        insert into classes (code, label, academic_year) values (:code, :label, :academic_year)
        returning id
        """),
        body.model_dump(),
    )
    return {"id": str(class_id), **body.model_dump(), "name": body.label, "active": True}


class SubjectIn(BaseModel):
    code: str = Field(min_length=1, max_length=32, pattern=r"^[0-9A-Za-z._-]+$")
    name: str = Field(min_length=1, max_length=160)


@router.get("/subjects")
def list_subjects(_: AdminOrTeacher, db: DB) -> list[dict[str, Any]]:
    rows = db.execute(text("select id, code, name from subjects order by name")).mappings()
    return [{**r, "id": str(r["id"])} for r in rows]


@router.post("/subjects", status_code=201)
def create_subject(body: SubjectIn, _: Admin, db: DB) -> dict[str, Any]:
    subject_id = db.scalar(
        text("insert into subjects (code, name) values (:code, :name) returning id"),
        body.model_dump(),
    )
    return {"id": str(subject_id), **body.model_dump()}


# ---- Staf & penugasan mengajar ---------------------------------------------------


class StaffIn(BaseModel):
    username: str = Field(pattern=USERNAME)
    display_name: str = Field(min_length=1, max_length=160)
    # Satu akun boleh guru sekaligus pengawas (BR: tanpa akun kedua).
    roles: set[Literal["TEACHER", "PROCTOR"]] = Field(min_length=1)


@router.get("/staff")
def list_staff(_: Admin, db: DB) -> list[dict[str, Any]]:
    rows = db.execute(
        text("""
        select u.id, u.username, u.display_name, u.status, array_agg(r.role order by r.role) roles
        from users u join user_roles r on r.user_id = u.id
        group by u.id
        having bool_or(r.role in ('ADMIN', 'TEACHER', 'PROCTOR'))
        order by u.display_name
        """)
    ).mappings()
    return [{**r, "id": str(r["id"])} for r in rows]


@router.post("/staff", status_code=201)
def create_staff(body: StaffIn, admin: Admin, db: DB) -> dict[str, Any]:
    password = generate_password()
    user_id = db.scalar(
        text("""
        insert into users (username, password_hash, display_name)
        values (:username, :hash, :name) returning id
        """),
        {"username": body.username, "hash": hash_password(password), "name": body.display_name},
    )
    db.execute(
        text("insert into user_roles (user_id, role, assigned_by) values (:id, :role, :by)"),
        [{"id": user_id, "role": r, "by": admin.id} for r in sorted(body.roles)],
    )
    # Kata sandi sementara hanya muncul sekali di respons ini (no-store), tidak pernah di log.
    return {
        "id": str(user_id),
        "username": body.username,
        "display_name": body.display_name,
        "roles": sorted(body.roles),
        "temporary_password": password,
    }


class TeachingAssignmentIn(BaseModel):
    teacher_id: UUID
    class_id: UUID
    subject_id: UUID


@router.post("/teaching-assignments", status_code=201)
def create_teaching_assignment(body: TeachingAssignmentIn, _: Admin, db: DB) -> dict[str, Any]:
    refs = db.execute(
        text("""
        select exists (select 1 from user_roles where user_id = :teacher_id and role = 'TEACHER'),
               exists (select 1 from classes where id = :class_id),
               exists (select 1 from subjects where id = :subject_id)
        """),
        body.model_dump(),
    ).one()
    if not all(refs):
        raise ApiError(422, "INVALID_REFERENCE", "Guru, kelas, atau mapel tidak valid.")
    assignment_id = db.scalar(
        text("""
        insert into teaching_assignments (teacher_id, class_id, subject_id)
        values (:teacher_id, :class_id, :subject_id) returning id
        """),
        body.model_dump(),
    )
    return {"id": str(assignment_id), **body.model_dump(mode="json")}


# ---- Siswa ----------------------------------------------------------------------


class ReasonIn(BaseModel):
    reason: str = Field(min_length=10, max_length=500)


@router.post("/students/{student_id}/password-reset")
def reset_student_password(student_id: UUID, body: ReasonIn, _: Admin, db: DB) -> dict[str, str]:
    # TODO(M8): simpan body.reason ke audit_events.
    password = generate_password()
    updated = db.scalar(
        text("""
        update users set password_hash = :hash, password_changed_at = now()
        where id = :id
          and exists (select 1 from user_roles where user_id = :id and role = 'STUDENT')
        returning id
        """),
        {"id": student_id, "hash": hash_password(password)},
    )
    if updated is None:
        raise not_found()
    # Mencabut semua sesi: juga jalan keluar bila siswa terkunci "akun aktif di perangkat lain".
    db.execute(text("delete from sessions where user_id = :id"), {"id": student_id})
    return {"temporary_password": password}


def cell_text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # NIS yang diketik sebagai angka di Excel
    return "" if value is None else str(value).strip()


def check_zip(data: bytes) -> zipfile.ZipFile:
    """Cek arsip XLSX/DOCX tak tepercaya SEBELUM parser XML menyentuhnya."""
    try:
        zf = zipfile.ZipFile(BytesIO(data))
        infos = zf.infolist()
    except zipfile.BadZipFile as e:
        raise ApiError(422, "INVALID_FILE", "Berkas bukan XLSX/DOCX yang valid.") from e
    if (
        len(infos) > 200
        or sum(i.file_size for i in infos) > MAX_UNZIPPED
        or any(i.file_size > 100 * max(i.compress_size, 1) for i in infos)
    ):
        raise ApiError(413, "FILE_TOO_LARGE", "Isi berkas terlalu besar atau mencurigakan.")
    if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
        raise ApiError(422, "MACRO_NOT_ALLOWED", "Berkas berisi makro. Simpan ulang tanpa makro.")
    return zf


def _parse_student_xlsx(data: bytes) -> list[tuple[int, str, str, str]]:
    """Baris (nomor_baris, nis, nama, kelas). Batas ukuran dicek sebelum openpyxl menyentuh XML."""
    check_zip(data)
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = [cell_text(c).lower() for c in next(rows, ())[:3]]
        if header != ["nis", "nama", "kelas"]:
            raise ApiError(
                422, "INVALID_TEMPLATE", "Baris pertama harus berisi kolom: NIS, Nama, Kelas."
            )
        out: list[tuple[int, str, str, str]] = []
        for number, row in enumerate(rows, start=2):
            nis, name, kelas = ([cell_text(c) for c in row[:3]] + ["", "", ""])[:3]
            if not (nis or name or kelas):
                continue
            if len(out) >= MAX_ROWS:
                raise ApiError(413, "TOO_MANY_ROWS", f"Maksimal {MAX_ROWS} siswa per berkas.")
            out.append((number, nis, name, kelas))
        wb.close()
    except ApiError:
        raise
    except Exception as e:  # berkas tak tepercaya: semua kegagalan parse = berkas tidak valid
        raise ApiError(422, "INVALID_FILE", "Berkas XLSX tidak dapat dibaca.") from e
    if not out:
        raise ApiError(422, "EMPTY_FILE", "Berkas tidak berisi data siswa.")
    return out


@router.post("/students/import/preview")
def preview_student_import(file: UploadFile, admin: Admin, db: DB) -> dict[str, Any]:
    # def biasa (bukan async): query DB sinkron berjalan di threadpool, tidak memblok event loop.
    data = file.file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise ApiError(413, "FILE_TOO_LARGE", "Ukuran berkas maksimal 2 MB.")
    rows = _parse_student_xlsx(data)

    nis_count = Counter(nis.lower() for _, nis, _, _ in rows)
    existing = set(
        db.scalars(
            text("select lower(username) from users where lower(username) = any(:l)"),
            {"l": list(nis_count)},
        )
    )
    classes = {
        code.lower(): (class_id, code)
        for class_id, code in db.execute(text("select id, code from classes where active"))
    }
    errors: list[dict[str, Any]] = []
    valid: list[dict[str, Any]] = []
    for number, nis, name, kelas in rows:
        row_errors = []
        if not NIS_RE.fullmatch(nis):
            row_errors.append(("NIS", "Format NIS tidak valid."))
        elif nis_count[nis.lower()] > 1:
            row_errors.append(("NIS", "NIS muncul lebih dari sekali di berkas."))
        elif nis.lower() in existing:
            row_errors.append(("NIS", "NIS sudah terdaftar."))
        if not name or len(name) > 160:
            row_errors.append(("Nama", "Nama wajib diisi, maksimal 160 karakter."))
        if kelas.lower() not in classes:
            row_errors.append(("Kelas", "Kelas tidak ditemukan atau tidak aktif."))
        errors += [{"row_number": number, "field": f, "message": m} for f, m in row_errors]
        if not row_errors:
            class_id, code = classes[kelas.lower()]
            valid.append({"nis": nis, "name": name, "class_id": str(class_id), "class": code})

    # Preview lama yang tak pernah di-commit ikut dibersihkan (berisi data siswa).
    db.execute(
        text(
            "delete from import_jobs where status = 'PREVIEW' "
            "and created_at < now() - interval '1 hour'"
        )
    )
    job_id = db.scalar(
        text("""
        insert into import_jobs (type, created_by, status, source_sha256, rows, errors)
        values ('STUDENT_XLSX', :by, 'PREVIEW', :sha, cast(:rows as jsonb), cast(:errors as jsonb))
        returning id
        """),
        {
            "by": admin.id,
            "sha": hashlib.sha256(data).hexdigest(),
            "rows": json.dumps(valid),
            "errors": json.dumps(errors),
        },
    )
    return {
        "preview_id": str(job_id),
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "errors": errors,
    }


class ImportCommitIn(BaseModel):
    preview_id: UUID
    confirm: Literal[True]


@router.post("/students/import/commit")
def commit_student_import(body: ImportCommitIn, admin: Admin, db: DB) -> dict[str, Any]:
    job = (
        db.execute(
            text("""
            select status, rows, errors from import_jobs
            where id = :id and type = 'STUDENT_XLSX' and created_by = :by
            for update
            """),
            {"id": body.preview_id, "by": admin.id},
        )
        .mappings()
        .first()
    )
    if job is None:
        raise not_found()
    if job["status"] != "PREVIEW":
        raise ApiError(409, "IMPORT_ALREADY_COMMITTED", "Import ini sudah disimpan.")
    if job["errors"]:
        raise ApiError(
            409, "IMPORT_HAS_ERRORS", "Perbaiki semua baris bermasalah lalu unggah ulang berkas."
        )

    rows: list[dict[str, Any]] = job["rows"]
    passwords = [generate_password() for _ in rows]
    params = [
        {**r, "id": uuid4(), "hash": h, "by": admin.id}
        for r, h in zip(rows, _hash_many(passwords), strict=True)
    ]
    # Satu transaksi: gagal di baris mana pun = tidak ada satu pun siswa tersimpan.
    try:
        db.execute(
            text(
                "insert into users (id, username, password_hash, display_name) "
                "values (:id, :nis, :hash, :name)"
            ),
            params,
        )
    except IntegrityError as e:
        raise ApiError(
            409, "IMPORT_CONFLICT", "Sebagian NIS baru saja terdaftar. Ulangi preview."
        ) from e
    db.execute(
        text("insert into user_roles (user_id, role, assigned_by) values (:id, 'STUDENT', :by)"),
        params,
    )
    db.execute(
        text("insert into class_memberships (class_id, student_id) values (:class_id, :id)"),
        params,
    )
    db.execute(
        text("update import_jobs set status = 'COMPLETED', rows = null where id = :id"),
        {"id": body.preview_id},
    )
    # Kredensial hanya dikembalikan sekali (no-store), tidak disimpan maupun di-log.
    return {
        "job_id": str(body.preview_id),
        "status": "COMPLETED",
        "credentials": [
            {"nis": r["nis"], "name": r["name"], "class": r["class"], "password": p}
            for r, p in zip(rows, passwords, strict=True)
        ],
    }
